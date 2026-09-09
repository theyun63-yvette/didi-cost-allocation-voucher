from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def display_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dict, list, tuple)):
        return str(value)
    return value


def rows_to_xlsx(sheet_rows: dict[str, list[dict[str, Any]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    if not sheet_rows:
        sheet_rows = {"结果": []}
    for sheet_name, rows in sheet_rows.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        headers: list[str] = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        if not headers:
            headers = ["说明"]
            rows = [{"说明": "无数据"}]
        ws.append(headers)
        for row in rows:
            ws.append([display_value(row.get(h, "")) for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, ws.max_column + 1):
            width = max(len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 200) + 1))
            ws.column_dimensions[ws.cell(1, col).column_letter].width = min(max(width + 2, 10), 42)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def file_digest(source: Any) -> str:
    if isinstance(source, bytes):
        data = source
    elif hasattr(source, "getvalue"):
        data = source.getvalue()
    else:
        with open(source, "rb") as f:
            data = f.read()
    return sha256(data).hexdigest()
