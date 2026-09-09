from __future__ import annotations

import io
from calendar import monthrange
from copy import copy
from datetime import datetime
from decimal import Decimal
from typing import Any
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from models.schemas import VoucherLine


def _load_template(source: Any):
    if isinstance(source, (str, Path)):
        return load_workbook(source)
    data = source if isinstance(source, bytes) else source.getvalue()
    return load_workbook(io.BytesIO(data))


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def generate_voucher(template_source: Any, template_meta: dict, target_month: str, lines: list[VoucherLine], config: dict) -> bytes:
    wb = _load_template(template_source)
    ws = wb[template_meta["sheet"]]
    start = template_meta["data_start_row"]
    original_max = ws.max_row
    needed_end = start + len(lines) - 1
    if needed_end > original_max:
        for row in range(original_max + 1, needed_end + 1):
            _copy_row_style(ws, start, row)
    for row in range(start, max(original_max, needed_end) + 1):
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).value = None
    year, month = [int(x) for x in target_month.split("-")]
    month_end = datetime(year, month, monthrange(year, month)[1])
    summary = f"计提{month}月{config['constants']['supplier_name']}交通费"
    for i, line in enumerate(lines):
        row = start + i
        if row > original_max:
            _copy_row_style(ws, start, row)
        if i == 0:
            ws.cell(row, 1).value = config["constants"]["book_code"]
            ws.cell(row, 2).value = month_end
            ws.cell(row, 3).value = month_end
            ws.cell(row, 4).value = config["constants"]["voucher_type_code"]
            ws.cell(row, 5).value = config["constants"]["voucher_type_name"]
            ws.cell(row, 6).value = None
        ws.cell(row, 7).value = summary
        ws.cell(row, 8).value = line.account_code
        ws.cell(row, 9).value = line.account_name
        ws.cell(row, 10).value = float(line.debit) if line.debit else None
        ws.cell(row, 11).value = float(line.credit) if line.credit else None
        ws.cell(row, 12).value = config["constants"]["currency_code"]
        ws.cell(row, 13).value = config["constants"]["currency_name"]
        if line.line_type == "费用":
            ws.cell(row, 14).value = "部门"
            ws.cell(row, 15).value = line.department_code
            ws.cell(row, 16).value = line.department_name
            ws.cell(row, 17).value = "项目"
            ws.cell(row, 18).value = line.project_code
            ws.cell(row, 19).value = line.project_name
        elif line.line_type == "应付账款":
            ws.cell(row, 14).value = line.dimension_type
            ws.cell(row, 15).value = line.dimension_code
            ws.cell(row, 16).value = line.dimension_name
        ws.cell(row, 10).number_format = "0.00"
        ws.cell(row, 11).number_format = "0.00"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
