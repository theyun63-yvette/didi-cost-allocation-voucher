import io
from openpyxl import Workbook
from parsers.didi_parser import parse_didi


def make_book(
    paid=103,
    refund=0,
    period="本期",
    settlement="2026-07-01 至 2026-07-31",
    order_id="O1",
    *,
    status="已完成",
    refund_time="",
    enterprise_payable=None,
):
    wb = Workbook()
    summary = wb.active
    summary.title = "汇总数据"
    summary.append(["结算周期", settlement])
    ws = wb.create_sheet("用车订单")
    ws.append(["说明"])
    ws.append(["所属账期", "企业订单号", "下单人姓名", "下单时间", "订单状态", "企业实付金额", "企业实退金额", "企业应付金额", "退款时间"])
    payable = paid - refund if enterprise_payable is None else enterprise_payable
    ws.append([period, order_id, "张三", "2026-07-05 10:00:00", status, paid, refund, payable, refund_time])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_refund_without_refund_time_blocks(config):
    result = parse_didi(make_book(refund=1, status="部分退款"), "2026-07", config)
    assert any(x.code == "E013" and x.blocking for x in result.issues)


def test_same_month_partial_refund_uses_net_amount(config):
    result = parse_didi(make_book(paid=103, refund=1, status="部分退款", refund_time="2026-07-08 10:00:00"), "2026-07", config)
    order = result.records[0]
    assert not [x for x in result.issues if x.blocking]
    assert order.valid
    assert order.original_enterprise_paid == 103
    assert order.enterprise_paid == 102
    assert order.refund_amount == 1
    assert order.enterprise_payable == 102
    assert order.refund_handling == "同月部分退款按净额计算"
    assert any(x.code == "W008" and not x.blocking for x in result.issues)


def test_same_month_full_refund_is_excluded_without_blocking(config):
    result = parse_didi(make_book(paid=103, refund=103, status="全额退款", refund_time="2026-07-08 10:00:00"), "2026-07", config)
    order = result.records[0]
    assert not [x for x in result.issues if x.blocking]
    assert not order.valid
    assert order.enterprise_paid == 0
    assert order.refund_handling == "同月全额退款排除"
    assert any(x.code == "W009" and not x.blocking for x in result.issues)


def test_cross_month_refund_blocks(config):
    result = parse_didi(make_book(refund=1, status="部分退款", refund_time="2026-08-01 10:00:00"), "2026-07", config)
    assert any(x.code == "E013" and x.blocking and "跨月退款" in x.reason for x in result.issues)


def test_refund_greater_than_paid_blocks(config):
    result = parse_didi(make_book(paid=100, refund=101, status="部分退款", refund_time="2026-07-08 10:00:00"), "2026-07", config)
    assert any(x.code == "E013" and x.blocking and "大于企业实付" in x.reason for x in result.issues)


def test_refund_net_must_match_enterprise_payable(config):
    result = parse_didi(make_book(paid=103, refund=1, status="部分退款", refund_time="2026-07-08 10:00:00", enterprise_payable=100), "2026-07", config)
    assert any(x.code == "E013" and x.blocking and "企业应付金额不一致" in x.reason for x in result.issues)


def test_negative_order_blocks(config):
    result = parse_didi(make_book(paid=-1), "2026-07", config)
    assert any(x.code == "E013" and x.blocking for x in result.issues)


def test_current_period_uses_summary_settlement_month(config):
    result = parse_didi(make_book(), "2026-07", config)
    assert not [x for x in result.issues if x.blocking]
    assert result.records[0].billing_period == "本期"
    assert result.records[0].billing_month == "2026-07"
    assert result.metadata["settlement_month"] == "2026-07"


def test_out_of_month_order_is_excluded_with_log(config):
    result = parse_didi(make_book(period="2026-06"), "2026-07", config)
    assert len(result.records) == 1
    assert not result.records[0].valid
    assert any(x.code == "W006" and not x.blocking for x in result.issues)


def test_zero_paid_order_is_excluded_with_log(config):
    result = parse_didi(make_book(paid=0), "2026-07", config)
    assert not result.records[0].valid
    assert any(x.code == "W004" and not x.blocking for x in result.issues)


def test_duplicate_order_id_blocks(config):
    wb = Workbook()
    summary = wb.active
    summary.title = "汇总数据"
    summary.append(["结算周期", "2026-07-01 至 2026-07-31"])
    ws = wb.create_sheet("用车订单")
    ws.append(["所属账期", "企业订单号", "下单人姓名", "企业实付金额", "企业实退金额"])
    ws.append(["本期", "O1", "张三", 103, 0])
    ws.append(["本期", "O1", "张三", 103, 0])
    b = io.BytesIO(); wb.save(b)
    result = parse_didi(b.getvalue(), "2026-07", config)
    assert any(x.code == "E011" and x.blocking for x in result.issues)
    assert not any(x.valid for x in result.records)
