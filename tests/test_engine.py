from __future__ import annotations

import io
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from services.engine import finalize_voucher, process_calculation, process_files
from conftest import make_template_bytes


def book_bytes(wb):
    b=io.BytesIO(); wb.save(b); return b.getvalue()


def personnel_book():
    wb=Workbook(); ws=wb.active; ws.title="人员"
    ws.append(["员工编号","姓名","部门名称"]); ws.append(["001","张三","开发"])
    return book_bytes(wb)


def allocation_book():
    wb=Workbook(); ws=wb.active; ws.title="工时分摊-2026.7"
    ws.append(["研发人员工时统计表"])
    ws.append(["序号","姓名","岗位","Coding","Prediction","AI4SCI","Multimodel","Cowork","Trading","Evaluation","非研发占比"])
    ws.append([1,"张三","开发",0.8,0,0,0,0,0,0,0.2])
    return book_bytes(wb)


def didi_book():
    wb=Workbook(); summary=wb.active; summary.title="汇总数据"
    summary.append(["结算周期","2026-07-01 至 2026-07-31"])
    ws=wb.create_sheet("用车订单")
    ws.append(["说明"]); ws.append(["所属账期","企业订单号","下单人姓名","支付时间","企业实付金额","企业实退金额"])
    ws.append(["本期","O1","张三","2026-07-10 12:00:00",103,0])
    return book_bytes(wb)


def test_end_to_end_generates_four_workbooks(config):
    result=process_files(personnel_book(),allocation_book(),didi_book(),make_template_bytes(),"2026-07",config)
    assert result.can_generate_voucher, [x.as_dict() for x in result.blocking_issues]
    assert set(result.output_files)=={"处理明细.xlsx","分摊汇总.xlsx","异常清单.xlsx","月度凭证.xlsx"}
    voucher=load_workbook(io.BytesIO(result.output_files["月度凭证.xlsx"]),data_only=True)
    ws=voucher["凭证导入"]
    assert ws["F5"].value is None
    assert ws["B5"].value.strftime("%Y-%m-%d")=="2026-07-31"
    assert sum((ws.cell(r,10).value or 0) for r in range(5,ws.max_row+1))==103
    assert sum((ws.cell(r,11).value or 0) for r in range(5,ws.max_row+1))==103


def test_two_stage_flow_waits_for_manual_confirmation(config):
    calculation = process_calculation(personnel_book(), allocation_book(), didi_book(), "2026-07", config)
    assert not calculation.blocking_issues
    assert set(calculation.output_files) == {"处理明细.xlsx", "分摊汇总.xlsx", "异常清单.xlsx"}
    assert calculation.voucher_lines == []
    assert "tax_amount" not in calculation.metrics

    unconfirmed = finalize_voucher(calculation, make_template_bytes(), config, confirmed=False)
    assert any(issue.code == "E020" and issue.blocking for issue in unconfirmed.issues)
    assert "月度凭证.xlsx" not in unconfirmed.output_files


def didi_partial_refund_book():
    wb=Workbook(); summary=wb.active; summary.title="汇总数据"
    summary.append(["结算周期","2026-07-01 至 2026-07-31"])
    ws=wb.create_sheet("用车订单")
    ws.append(["说明"])
    ws.append(["所属账期","企业订单号","下单人姓名","支付时间","订单状态","企业实付金额","企业实退金额","企业应付金额","退款时间"])
    ws.append(["本期","O1","张三","2026-07-10 12:00:00","部分退款",103,1,102,"2026-07-10 13:00:00"])
    return book_bytes(wb)


def test_same_month_partial_refund_flows_through_allocation_and_voucher(config):
    result = process_files(personnel_book(), allocation_book(), didi_partial_refund_book(), make_template_bytes(), "2026-07", config)
    assert result.can_generate_voucher, [x.as_dict() for x in result.blocking_issues]
    assert result.metrics["gross_amount"] == Decimal("102")
    assert result.metrics["monthly_net_amount"] == Decimal("99.03")
    assert result.metrics["tax_amount"] == Decimal("2.97")
    assert result.metrics["partial_refund_order_count"] == 1
    row = result.order_detail_rows[0]
    assert row["企业实付金额（原始）"] == Decimal("103")
    assert row["企业实退金额"] == Decimal("1")
    assert row["有效价税合计"] == Decimal("102")
    assert row["退款处理方式"] == "同月部分退款按净额计算"
    assert any(x.code == "W008" and not x.blocking for x in result.issues)
