import io
from decimal import Decimal
from openpyxl import load_workbook
from models.schemas import VoucherLine
from parsers.voucher_template_parser import parse_voucher_template
from services.voucher import generate_voucher
from conftest import make_template_bytes


def lines():
    return [
        VoucherLine("费用","6602.07","日常交通费",debit=Decimal("100"),department_code="008",department_name="职能",project_code="00009",project_name="非研发"),
        VoucherLine("税额","2221.04","待认证进项税额",debit=Decimal("3")),
        VoucherLine("应付账款","2202","应付账款",credit=Decimal("103"),dimension_type="供应商",dimension_code="SC00020",dimension_name="滴滴出行科技有限公司"),
    ]


def test_voucher_number_blank(config):
    template=make_template_bytes(); meta=parse_voucher_template(template).metadata
    data=generate_voucher(template,meta,"2026-07",lines(),config)
    ws=load_workbook(io.BytesIO(data),data_only=True)["凭证导入"]
    assert ws["F5"].value is None


def test_month_generates_one_voucher_header(config):
    template=make_template_bytes(); meta=parse_voucher_template(template).metadata
    data=generate_voucher(template,meta,"2026-07",lines(),config)
    ws=load_workbook(io.BytesIO(data),data_only=True)["凭证导入"]
    assert sum(1 for r in range(5,8) if ws.cell(r,1).value)==1


def test_template_hidden_sheet_and_structure_preserved(config):
    template=make_template_bytes(hidden=True); parsed=parse_voucher_template(template)
    data=generate_voucher(template,parsed.metadata,"2026-07",lines(),config)
    wb=load_workbook(io.BytesIO(data),data_only=True)
    assert wb["下拉选项"].sheet_state=="hidden"
    assert wb["下拉选项"]["A1"].value=="保留"
    assert wb["凭证导入"]["A3"].value=="book.number"
