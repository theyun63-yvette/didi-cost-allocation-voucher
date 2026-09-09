from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook

from parsers.allocation_parser import parse_allocation
from parsers.personnel_parser import parse_personnel
from services.engine import process_files
from conftest import make_template_bytes

PROJECTS = ["Coding", "Prediction", "AI4SCI", "Multimodel", "Cowork", "Trading", "Evaluation"]


def book_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def personnel_book(rows=None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员"
    sheet.append(["员工编号", "姓名", "部门名称"])
    for row in rows or [["001", "张三", "开发"]]:
        sheet.append(row)
    return book_bytes(workbook)


def allocation_book(*, project_values=None, non_rnd=Decimal("0.2"), extra_header=None) -> bytes:
    values = {name: Decimal("0") for name in PROJECTS}
    values.update({"Coding": Decimal("0.8")} if project_values is None else project_values)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工时分摊-2026.7"
    sheet.append(["研发人员工时统计表"])
    headers = ["序号", "姓名", "岗位", *PROJECTS, "非研发占比"]
    if extra_header:
        headers.append(extra_header)
    sheet.append(headers)
    row = [1, "张三", "开发", *[values[name] for name in PROJECTS], non_rnd]
    if extra_header:
        row.append(Decimal("0.1"))
    sheet.append(row)
    return book_bytes(workbook)


def didi_book(*, name="张三", paid=Decimal("103"), period="本期") -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总数据"
    summary.append(["结算周期", "2026-07-01 至 2026-07-31"])
    sheet = workbook.create_sheet("用车订单")
    sheet.append(["所属账期", "企业订单号", "下单人姓名", "企业实付金额", "企业实退金额"])
    sheet.append([period, "O1", name, paid, 0])
    return book_bytes(workbook)


def test_blank_rnd_ratio_becomes_zero_and_logs(config):
    result = parse_allocation(allocation_book(project_values={"Coding": Decimal("0.8"), "Trading": None}), "2026-07", config)
    record = result.records[0]
    assert record.project_ratios["Trading"] == Decimal("0")
    assert "Trading" in record.blank_projects
    assert any(issue.code == "W003" and issue.employee_name == "张三" and "Trading" in issue.reason for issue in result.issues)


def test_all_seven_projects_and_non_rnd_sum_to_one(config):
    values = {name: Decimal("0.1") for name in PROJECTS}
    result = parse_allocation(allocation_book(project_values=values, non_rnd=Decimal("0.3")), "2026-07", config)
    assert len(result.records) == 1
    assert result.records[0].valid
    assert sum(result.records[0].project_ratios.values(), Decimal("0")) + result.records[0].non_rnd_ratio == Decimal("1.0")
    assert not [issue for issue in result.issues if issue.blocking]


def test_blank_non_rnd_is_zero_with_log_when_timesheet_record_is_matched(config):
    result = process_files(personnel_book(), allocation_book(project_values={"Coding": Decimal("1")}, non_rnd=None), didi_book(), make_template_bytes(), "2026-07", config)
    assert not result.blocking_issues
    assert any(issue.code == "W007" and not issue.blocking for issue in result.issues)
    assert len(result.order_detail_rows) == 1
    assert result.order_detail_rows[0]["非研发金额"] == Decimal("0")
    assert "月度凭证.xlsx" in result.output_files


def test_same_name_same_accounting_role_is_deduped(config):
    result = parse_personnel(personnel_book([["001", "张三", "开发"], ["002", "张三", "开发"]]), config)
    assert len(result.records) == 1
    assert any(issue.code == "W002" and not issue.blocking for issue in result.issues)
    assert not [issue for issue in result.issues if issue.blocking]


def test_same_name_different_accounting_roles_blocks(config):
    result = parse_personnel(personnel_book([["001", "张三", "开发"], ["002", "张三", "算法"]]), config)
    assert any(issue.code == "E004" and issue.blocking for issue in result.issues)


def test_order_detail_and_summary_are_balanced(config):
    result = process_files(personnel_book(), allocation_book(), didi_book(), make_template_bytes(), "2026-07", config)
    assert result.can_generate_voucher
    row = result.order_detail_rows[0]
    assert row["单笔不含税金额"] == Decimal("100.00")
    assert row["Coding金额"] == Decimal("80.000")
    assert row["非研发金额"] == Decimal("20.000")
    assert row["分摊金额合计"] == Decimal("100.000")
    assert row["单笔分摊差额"] == Decimal("0.000")
    assert result.metrics["order_detail_allocation_total"] == Decimal("100.000")
    assert result.metrics["allocation_summary_total"] == Decimal("100.000")
    assert result.summary_rows[-1]["会计归属类别/岗位"] == "合计"


def test_order_net_uses_round_half_up(config):
    result = process_files(personnel_book(), allocation_book(project_values={}, non_rnd=Decimal("1")), didi_book(paid=Decimal("1.03515")), make_template_bytes(), "2026-07", config)
    assert result.order_detail_rows[0]["单笔不含税金额"] == Decimal("1.01")


def test_same_inputs_produce_identical_calculation_results(config):
    args = (personnel_book(), allocation_book(), didi_book(), make_template_bytes(), "2026-07", config)
    first = process_files(*args)
    second = process_files(*args)
    assert first.order_detail_rows == second.order_detail_rows
    assert first.summary_rows == second.summary_rows
    assert [line.as_dict() for line in first.voucher_lines] == [line.as_dict() for line in second.voucher_lines]


def test_unknown_timesheet_project_blocks(config):
    result = parse_allocation(allocation_book(extra_header="NewProject"), "2026-07", config)
    assert any(issue.code == "E009" and issue.blocking and issue.original_value == "NewProject" for issue in result.issues)


def test_missing_personnel_blocks(config):
    result = process_files(personnel_book(), allocation_book(), didi_book(name="不存在人员"), make_template_bytes(), "2026-07", config)
    assert any(issue.code == "E003" and issue.blocking for issue in result.issues)


def test_name_whitespace_is_normalized_and_logged(config):
    result = process_files(personnel_book([["001", " 张三　", "开发"]]), allocation_book(), didi_book(name=" 张三　"), make_template_bytes(), "2026-07", config)
    assert result.order_detail_rows[0]["标准化姓名"] == "张三"
    assert any(issue.code == "W001" for issue in result.issues)


def allocation_book_without_evaluation(month: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"工时分摊-2026.{month}"
    headers = ["序号", "姓名", "岗位", "Coding", "Prediction", "AI4SCI", "Multimodel", "Cowork", "Trading", "非研发占比"]
    sheet.append(["研发人员工时统计表"])
    sheet.append(headers)
    sheet.append([1, "张三", "开发", Decimal("0.8"), 0, 0, 0, 0, 0, Decimal("0.2")])
    return book_bytes(workbook)


def test_project_missing_before_effective_month_is_zero_with_log(config):
    result = parse_allocation(allocation_book_without_evaluation(5), "2026-05", config)
    assert not [issue for issue in result.issues if issue.blocking]
    assert result.records[0].project_ratios["Evaluation"] == Decimal("0")
    assert result.metadata["historical_inactive_projects"] == ["Evaluation"]
    assert any(issue.code == "W010" and not issue.blocking and issue.original_value == "Evaluation" for issue in result.issues)


def test_project_missing_from_effective_month_blocks(config):
    result = parse_allocation(allocation_book_without_evaluation(6), "2026-06", config)
    assert any(issue.code == "E002" and issue.blocking and issue.original_value == "Evaluation" for issue in result.issues)
