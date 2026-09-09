from __future__ import annotations

import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u3000", " ").strip()
    return re.sub(r"\s+", " ", text)


def normalize_role(value: Any, aliases: dict[str, str] | None = None) -> str:
    text = normalize_text(value)
    return (aliases or {}).get(text, text)


def as_decimal(value: Any, *, blank_zero: bool = False) -> Decimal:
    if value is None or normalize_text(value) == "":
        if blank_zero:
            return Decimal("0")
        raise InvalidOperation("blank")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise InvalidOperation("boolean")
    text = normalize_text(value).replace(",", "")
    if text.endswith("%"):
        return Decimal(text[:-1]) / Decimal("100")
    return Decimal(text)


def source_name(source: Any, fallback: str) -> str:
    name = getattr(source, "name", None)
    if name:
        return Path(str(name)).name
    if isinstance(source, (str, Path)):
        return Path(source).name
    return fallback


def load_xlsx(source: Any, *, data_only: bool = True):
    if isinstance(source, (str, Path)):
        return load_workbook(source, data_only=data_only)
    if hasattr(source, "getvalue"):
        data = source.getvalue()
    elif hasattr(source, "read"):
        pos = source.tell() if hasattr(source, "tell") else None
        data = source.read()
        if pos is not None and hasattr(source, "seek"):
            source.seek(pos)
    elif isinstance(source, bytes):
        data = source
    else:
        raise TypeError(f"Unsupported Excel source: {type(source)!r}")
    return load_workbook(io.BytesIO(data), data_only=data_only)


def resolve_header(headers: list[Any], aliases: list[str], field_name: str, required: bool = True) -> tuple[int | None, str | None]:
    normalized = [normalize_text(x) for x in headers]
    alias_set = {normalize_text(x) for x in aliases}
    matches = [(i, h) for i, h in enumerate(normalized) if h in alias_set]
    if len(matches) > 1:
        raise ValueError(f"字段“{field_name}”匹配到多个表头: {[x[1] for x in matches]}")
    if not matches:
        if required:
            raise KeyError(field_name)
        return None, None
    return matches[0]


def find_header_row(ws, required_alias_groups: list[list[str]], max_rows: int = 30) -> int | None:
    candidates = []
    for row in range(1, min(ws.max_row, max_rows) + 1):
        headers = [normalize_text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        ok = all(any(normalize_text(alias) in headers for alias in group) for group in required_alias_groups)
        if ok:
            candidates.append(row)
    return candidates[0] if len(candidates) == 1 else None
